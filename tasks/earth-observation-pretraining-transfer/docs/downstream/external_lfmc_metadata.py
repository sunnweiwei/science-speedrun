"""Select geographic evaluation candidates from metadata only, never LFMC values."""

import argparse
from collections import Counter, defaultdict
from datetime import datetime, date
import hashlib
import json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.spatial import cKDTree

p = argparse.ArgumentParser()
p.add_argument("--workbook", type=Path, required=True)
p.add_argument("--policy", type=Path, required=True)
p.add_argument("--pretrain-coordinates", type=Path, nargs="+", required=True)
p.add_argument("--output", type=Path, required=True)
a = p.parse_args()
if a.output.exists():
    raise ValueError("Preserve existing candidate selection")
policy = json.loads(a.policy.read_text())
with a.workbook.open("rb") as f:
    md5 = hashlib.file_digest(f, "md5").hexdigest()
if md5 != policy["file_md5"]:
    raise ValueError("Wrong pinned workbook")
coords = np.concatenate([np.load(p) for p in a.pretrain_coordinates])
tree = cKDTree(coords)
book = load_workbook(a.workbook, read_only=True, data_only=True)
iterator = book["LFMC data"].iter_rows(values_only=True)
header = next(iterator)
needed = [
    "Latitude (WGS84, EPSG:4326)",
    "Longitude (WGS84, EPSG:4326)",
    "Sampling date (YYYYMMDD)",
    "Species functional type",
    "Country",
]
columns = {name: header.index(name) for name in needed}
assert "LFMC value (%)" not in columns
sites = {}
counters = Counter()
bounds = policy["coordinate_bounds"]
for row_number, row in enumerate(iterator, start=2):
    counters["source_rows"] += 1
    # Only these metadata columns are selected. No target column is indexed.
    values = {name: row[index] for name, index in columns.items()}
    if (
        str(values["Species functional type"]).lower().strip()
        not in policy["woody_functional_types"]
    ):
        continue
    try:
        lat = float(values[needed[0]])
        lon = float(values[needed[1]])
        raw = values[needed[2]]
        when = (
            raw.date()
            if isinstance(raw, datetime)
            else raw
            if isinstance(raw, date)
            else datetime.strptime(str(int(raw)), "%Y%m%d").date()
        )
    except (TypeError, ValueError, OverflowError):
        counters["invalid_metadata"] += 1
        continue
    if not np.isfinite([lat, lon]).all() or not (
        bounds["min_lon"] <= lon <= bounds["max_lon"]
        and bounds["min_lat"] <= lat <= bounds["max_lat"]
    ):
        continue
    if not (
        policy["eligible_dates"][0] <= when.isoformat() <= policy["eligible_dates"][1]
    ):
        continue
    counters["geography_date_woody_rows"] += 1
    if tree.query([lat, lon])[0] < 0.1:
        counters["pretraining_buffer_excluded_rows"] += 1
        continue
    key = f"{lat:.6f},{lon:.6f}"
    site = sites.setdefault(
        key,
        {
            "latitude": lat,
            "longitude": lon,
            "country": str(values["Country"]),
            "dates": defaultdict(list),
        },
    )
    site["dates"][when.isoformat()].append(row_number)
book.close()
sites = {
    k: v
    for k, v in sites.items()
    if len(v["dates"]) >= policy["minimum_distinct_dates_per_site"]
}


def order(value):
    return hashlib.sha256(("external-lfmc-v1|" + value).encode()).hexdigest()


selected = []
for key in sorted(sites, key=order)[: policy["maximum_sites"]]:
    site = sites[key]
    dates = sorted(site["dates"], key=lambda d: order(key + "|" + d))[
        : policy["maximum_dates_per_site"]
    ]
    selected.append(
        {
            "site_id": key,
            "latitude": site["latitude"],
            "longitude": site["longitude"],
            "country": site["country"],
            "observations": [
                {"date": d, "source_excel_rows": site["dates"][d]}
                for d in sorted(dates)
            ],
        }
    )
result = {
    "policy_sha256": hashlib.sha256(a.policy.read_bytes()).hexdigest(),
    "source_md5": md5,
    "selected_metadata_columns": needed,
    "target_column_selected": False,
    "target_statistics_computed": False,
    "model_predictions_used": False,
    "source_counters": dict(counters),
    "eligible_sites": len(sites),
    "selected_sites": len(selected),
    "selected_site_dates": sum(len(s["observations"]) for s in selected),
    "selected_sites_by_country": dict(Counter(s["country"] for s in selected)),
    "pretraining_coordinate_sources": {
        str(p): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in a.pretrain_coordinates
    },
    "sites": selected,
    "status": "Metadata candidates only; imagery and source-label processing audit pending, no scores computed",
}
a.output.parent.mkdir(parents=True, exist_ok=True)
a.output.write_text(json.dumps(result, indent=2) + "\n")
print(json.dumps({k: v for k, v in result.items() if k != "sites"}, indent=2))
